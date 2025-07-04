import gradio as gr
from sentence_transformers import SentenceTransformer, util
from pdfminer.high_level import extract_text
from docx import Document
import os
import re
import pandas as pd
from fpdf import FPDF
import google.generativeai as genai

# 🔐 Configure Gemini API Key
genai.configure(api_key="Your_api_key_here")  # Replace with your actual Gemini API key

# Load embedding model
model = SentenceTransformer('all-MiniLM-L6-v2')

def extract_resume_text(file_path):
    ext = os.path.splitext(file_path)[-1].lower()
    if ext == '.pdf':
        return extract_text(file_path)
    elif ext == '.docx':
        doc = Document(file_path)
        return "\n".join([para.text for para in doc.paragraphs])
    return "Unsupported file format"

def clean_text(text):
    text = re.sub(r'\s+', ' ', text)
    text = re.sub(r'[^\x00-\x7F]+', '', text)
    return text.strip().lower()

def extract_keywords(text):
    words = re.findall(r'\b[a-zA-Z]{3,}\b', text.lower())
    stopwords = set([
        "with", "the", "and", "for", "you", "are", "our", "that", "from", "will", "have",
        "your", "job", "role", "responsibilities", "requirements", "preferred", "must"
    ])
    return list(set([word for word in words if word not in stopwords]))

def keyword_matcher(resume_text, jd_text):
    jd_keywords = extract_keywords(jd_text)
    matched = [kw for kw in jd_keywords if kw in resume_text]
    missing = [kw for kw in jd_keywords if kw not in resume_text]
    return matched, missing

# ✅ Gemini GPT Feedback
def gpt_feedback(resume_text, jd_text):
    prompt = f"""
You are an expert recruiter. Evaluate the following resume for the given job description.
Give a match score (0-100), 2 strengths, and 2 improvement suggestions. dont include ** this symbols in the output 

Job Description:
{jd_text}

Resume:
{resume_text}
"""
    try:
        model = genai.GenerativeModel('gemini-2.0-flash')
        response = model.generate_content(prompt)
        return response.text
    except Exception as e:
        return f"⚠️ Gemini feedback failed: {e}"

def evaluate_resume(resume_file, job_description, use_gpt=False):
    if not resume_file or not job_description:
        return "Missing input", "", "", "", ""

    resume_text = extract_resume_text(resume_file.name)
    resume_text = clean_text(resume_text)
    job_description = clean_text(job_description)

    embeddings = model.encode([resume_text, job_description])
    similarity = util.cos_sim(embeddings[0], embeddings[1]).item()
    score = round(similarity * 100, 2)

    explanation = "✅ Excellent fit!" if score > 80 else (
                  "👍 Good match." if score > 60 else (
                  "⚠️ Fair match." if score > 40 else "❌ Weak match."))

    matched, missing = keyword_matcher(resume_text, job_description)
    gpt_output = gpt_feedback(resume_text, job_description) if use_gpt else ""

    return f"{score}%", explanation, ", ".join(matched), ", ".join(missing), gpt_output

def rank_resumes(resume_files, job_description):
    job_description = clean_text(job_description)
    results = []

    for file in resume_files:
        try:
            resume_text = extract_resume_text(file.name)
            resume_text = clean_text(resume_text)
            embeddings = model.encode([resume_text, job_description])
            similarity = util.cos_sim(embeddings[0], embeddings[1]).item()
            score = round(similarity * 100, 2)
            results.append((os.path.basename(file.name), score))
        except Exception as e:
            results.append((os.path.basename(file.name), f"Error: {e}"))

    sorted_results = sorted(results, key=lambda x: (x[1] if isinstance(x[1], (int, float)) else -1), reverse=True)
    return sorted_results

def generate_csv_report(results):
    df = pd.DataFrame(results, columns=["Resume", "Score"])
    csv_path = "resume_ranking_report.csv"
    df.to_csv(csv_path, index=False)
    return csv_path

def generate_pdf_report(results):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt="Resume Ranking Report", ln=True, align='C')
    pdf.ln(10)

    pdf.set_font("Arial", size=10)
    pdf.cell(80, 10, "Resume", 1)
    pdf.cell(30, 10, "Score", 1)
    pdf.ln()

    for resume, score in results:
        pdf.cell(80, 10, str(resume), 1)
        pdf.cell(30, 10, str(score), 1)
        pdf.ln()

    pdf_path = "resume_ranking_report.pdf"
    pdf.output(pdf_path)
    return pdf_path

# 🚀 Gradio UI
with gr.Blocks() as demo:
    gr.Markdown("## 🤖 AI Resume & Job Match Evaluator")

    with gr.Tab("🔍 Evaluate One Resume"):
        resume_input = gr.File(label="📄 Upload Resume", file_types=[".pdf", ".docx"])
        jd_input = gr.Textbox(label="📌 Job Description", lines=10)
        use_gpt = gr.Checkbox(label="🔮 Use Gemini for Smart Feedback")

        with gr.Row():
            score_output = gr.Textbox(label="Score")
            explanation_output = gr.Textbox(label="Feedback")

        with gr.Accordion("🧠 Skill Match", open=False):
            matched_output = gr.Textbox(label="✅ Matched Skills")
            missing_output = gr.Textbox(label="❌ Missing Skills")

        gpt_output = gr.Textbox(label="🧠 Gemini Feedback", lines=6)

        eval_btn = gr.Button("Evaluate")
        eval_btn.click(
            fn=evaluate_resume,
            inputs=[resume_input, jd_input, use_gpt],
            outputs=[score_output, explanation_output, matched_output, missing_output, gpt_output]
        )

    with gr.Tab("📊 Rank Multiple Resumes"):
        resume_files = gr.File(label="Upload Multiple Resumes", file_types=[".pdf", ".docx"], file_count="multiple")
        multi_jd_input = gr.Textbox(label="Job Description", lines=8)
        rank_btn = gr.Button("🏆 Rank Resumes")
        rank_output = gr.Textbox(label="Ranking Result", lines=10)
        download_csv = gr.File(label="Download CSV Report")
        download_pdf = gr.File(label="Download PDF Report")

        def rank_and_report(resume_files, jd_text):
            results = rank_resumes(resume_files, jd_text)
            display = "\n".join([f"{i+1}. {name} - {score}" for i, (name, score) in enumerate(results)])
            csv_path = generate_csv_report(results)
            pdf_path = generate_pdf_report(results)
            return display, csv_path, pdf_path

        rank_btn.click(fn=rank_and_report, inputs=[resume_files, multi_jd_input],
                       outputs=[rank_output, download_csv, download_pdf])

demo.launch()
